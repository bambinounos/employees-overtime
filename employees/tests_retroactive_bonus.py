"""
Pruebas de cobro retroactivo justo por errores administrativos (KPIs count_lt y Opción A con IPAC)
y soporte responsivo móvil en Dashboard y Admin.
"""
import datetime
import io
from datetime import date, time
from decimal import Decimal

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .exports import generar_planilla_xlsx
from .models import (
    BonusRule,
    CompanySettings,
    DescuentoRetroactivoBono,
    Employee,
    JobProfile,
    KPI,
    ManualKpiEntry,
    Salary,
    Task,
    TaskBoard,
    TaskList,
    WorkLog,
)
from .nomina import generar_recibo
from .report_pdf import generar_recibo_pdf


def _aware(year, month, day, hour=12, minute=0):
    tz = timezone.get_current_timezone()
    return timezone.make_aware(datetime.datetime.combine(date(year, month, day), time(hour, minute)), tz)


class RetroactiveBonusDeductionTests(TestCase):
    def setUp(self):
        CompanySettings.objects.update_or_create(
            pk=1, defaults={'calculation_basis': 'monthly', 'base_hours': Decimal('160.00')}
        )
        self.kpi_calidad = KPI.objects.create(
            name='Calidad Administrativa',
            measurement_type='count_lt',
            target_value=Decimal('2.00'),  # máximo 1 error permitido (< 2)
        )
        BonusRule.objects.create(kpi=self.kpi_calidad, bonus_amount=Decimal('30.00'))

        self.kpi_ipac = KPI.objects.create(
            name='IPAC',
            measurement_type='composite_ipac',
            target_value=Decimal('1.50'),
        )
        BonusRule.objects.create(kpi=self.kpi_ipac, bonus_amount=Decimal('50.00'))

        self.profile = JobProfile.objects.create(name='Administrativo')
        self.profile.kpis.add(self.kpi_calidad, self.kpi_ipac)

        self.user = User.objects.create_user('empleado1', 'emp1@example.com', 'pass1234')
        self.superuser = User.objects.create_superuser('admin1', 'admin1@example.com', 'pass1234')

        self.employee = Employee.objects.create(
            user=self.user,
            name='Empleado Prueba',
            email='emp1@example.com',
            hire_date=date(2026, 1, 1),
            profile=self.profile,
        )
        Salary.objects.create(
            employee=self.employee,
            base_amount=Decimal('500.00'),
            effective_date=date(2026, 1, 1),
        )
        # 160 horas en agosto (2026-08) y septiembre (2026-09)
        for m in (8, 9):
            for d in range(1, 21):
                WorkLog.objects.create(
                    employee=self.employee, date=date(2026, m, d), hours_worked=Decimal('8.00')
                )

        self.board = TaskBoard.objects.create(employee=self.employee, name='Tablero')
        self.list_done = TaskList.objects.create(board=self.board, name='Hecho', order=3)

    def _create_completed_tasks(self, year, month, count):
        for i in range(count):
            t = Task.objects.create(
                list=self.list_done,
                assigned_to=self.employee,
                title=f"Tarea {i+1}",
                order=i + 1,
                due_date=_aware(year, month, 20, 17, 0),
                status='completed',
            )
            Task.objects.filter(pk=t.pk).update(completed_at=_aware(year, month, 15, 10, 0))

    def test_retroactive_error_deducts_from_future_salary_without_polluting_current_month(self):
        """
        En agosto (2026-08) el empleado tuvo 1 error a tiempo (permitido: máximo 1) y cobró $30.
        En septiembre (2026-09) se descubre un 2.º error ocurrido en agosto (date=2026-08-18,
        registrado el 2026-09-10).
        - Agosto mantiene su registro histórico de lo cobrado ($30 de Calidad).
        - Septiembre NO se contamina con el error de agosto (gana sus $30 de septiembre).
        - En septiembre se aplica un DescuentoRetroactivoBono de -$30.00 por el bono de agosto.
        """
        # 1 error a tiempo en agosto -> aún cumple (< 2)
        e1 = ManualKpiEntry.objects.create(
            employee=self.employee,
            kpi=self.kpi_calidad,
            date=date(2026, 8, 10),
            value=Decimal('1'),
        )
        ManualKpiEntry.objects.filter(pk=e1.pk).update(created_at=_aware(2026, 8, 10, 14, 0))

        sal_aug = self.employee.calculate_salary(2026, 8)
        self.assertEqual(sal_aug['performance_bonus'], Decimal('30.00'))
        self.assertEqual(sal_aug['retroactive_bonus_deduction'], Decimal('0.00'))
        self.assertEqual(sal_aug['total_salary'], Decimal('530.00'))

        # En septiembre se detecta un 2.º error que ocurrió el 18 de agosto
        e2 = ManualKpiEntry.objects.create(
            employee=self.employee,
            kpi=self.kpi_calidad,
            date=date(2026, 8, 18),
            value=Decimal('1'),
            notes='Error detectado en auditoría de septiembre',
        )
        ManualKpiEntry.objects.filter(pk=e2.pk).update(created_at=_aware(2026, 9, 10, 11, 0))
        self.employee.sync_retroactive_bonus_deductions()

        # Verificar que el descuento retroactivo se asignó a 09/2026 por $30.00
        deds = DescuentoRetroactivoBono.objects.filter(
            employee=self.employee, origin_year=2026, origin_month=8, applied_year=2026, applied_month=9
        )
        self.assertEqual(deds.count(), 1)
        self.assertEqual(deds.first().amount, Decimal('30.00'))

        # Calcular sueldo de septiembre: gana $30 de bono en sep, pero se descuentan -$30 retroactivos de ago
        sal_sep = self.employee.calculate_salary(2026, 9)
        self.assertEqual(sal_sep['performance_bonus'], Decimal('30.00'))
        self.assertEqual(sal_sep['retroactive_bonus_deduction'], Decimal('30.00'))
        self.assertEqual(sal_sep['total_salary'], Decimal('500.00'))

        # Si se registra un 3.er error retroactivo de agosto, NO vuelve a descontar otros $30
        e3 = ManualKpiEntry.objects.create(
            employee=self.employee,
            kpi=self.kpi_calidad,
            date=date(2026, 8, 22),
            value=Decimal('1'),
        )
        ManualKpiEntry.objects.filter(pk=e3.pk).update(created_at=_aware(2026, 9, 15, 11, 0))
        sal_sep_after = self.employee.calculate_salary(2026, 9)
        self.assertEqual(sal_sep_after['retroactive_bonus_deduction'], Decimal('30.00'))

    def test_option_a_recalculates_both_count_lt_and_ipac_when_threshold_crossed(self):
        """
        Opción A: En agosto (21 días hábiles) el empleado completó 36 tareas a tiempo y tuvo
        1 error a tiempo -> IPAC = 36 * 0.90 / 21 = 1.54 (>= 1.50, gana $50) y Calidad = 1 (< 2, gana $30).
        Total bono cobrado en agosto = $80.00.
        En septiembre se descubre un 2.º error ocurrido en agosto:
        - Calidad en agosto pasa de 1 a 2 errores -> pierde $30.00.
        - IPAC en agosto pasa a 36 * 0.80 / 21 = 1.37 (< 1.50) -> pierde $50.00.
        - Total descuento retroactivo en septiembre = $80.00.
        """
        self._create_completed_tasks(2026, 8, 36)
        e1 = ManualKpiEntry.objects.create(
            employee=self.employee,
            kpi=self.kpi_calidad,
            date=date(2026, 8, 10),
            value=Decimal('1'),
        )
        ManualKpiEntry.objects.filter(pk=e1.pk).update(created_at=_aware(2026, 8, 10, 14, 0))

        sal_aug = self.employee.calculate_salary(2026, 8)
        self.assertEqual(sal_aug['performance_bonus'], Decimal('80.00'))

        # Hallazgo retroactivo en septiembre
        e2 = ManualKpiEntry.objects.create(
            employee=self.employee,
            kpi=self.kpi_calidad,
            date=date(2026, 8, 20),
            value=Decimal('1'),
        )
        ManualKpiEntry.objects.filter(pk=e2.pk).update(created_at=_aware(2026, 9, 5, 9, 0))

        sal_sep = self.employee.calculate_salary(2026, 9)
        self.assertEqual(sal_sep['retroactive_bonus_deduction'], Decimal('80.00'))
        self.assertEqual(len(sal_sep['retroactive_deductions']), 2)

        # Si se elimina el error retroactivo, el descuento desaparece automáticamente
        e2.delete()
        sal_sep_clean = self.employee.calculate_salary(2026, 9)
        self.assertEqual(sal_sep_clean['retroactive_bonus_deduction'], Decimal('0.00'))

    def test_pdf_and_xlsx_include_retroactive_deduction(self):
        e1 = ManualKpiEntry.objects.create(
            employee=self.employee, kpi=self.kpi_calidad, date=date(2026, 8, 5), value=Decimal('1')
        )
        ManualKpiEntry.objects.filter(pk=e1.pk).update(created_at=_aware(2026, 8, 5, 10, 0))
        self.employee.calculate_salary(2026, 8)

        e2 = ManualKpiEntry.objects.create(
            employee=self.employee, kpi=self.kpi_calidad, date=date(2026, 8, 15), value=Decimal('1')
        )
        ManualKpiEntry.objects.filter(pk=e2.pk).update(created_at=_aware(2026, 9, 12, 10, 0))

        recibo = generar_recibo(self.employee, 2026, 9, generado_por=self.superuser)
        self.assertEqual(recibo.datos['retroactive_bonus_deduction'], '30.00')
        self.assertEqual(recibo.total, Decimal('500.00'))

        pdf_bytes = generar_recibo_pdf(recibo)
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))

        xlsx_bytes, preliminar = generar_planilla_xlsx(2026, 9)
        self.assertFalse(preliminar)
        ws = load_workbook(io.BytesIO(xlsx_bytes)).active
        self.assertEqual(Decimal(str(ws.cell(row=8, column=10).value)), Decimal('30.00'))
        self.assertEqual(Decimal(str(ws.cell(row=8, column=11).value)), Decimal('500.00'))


class MobileResponsivenessTemplateTests(TestCase):
    def setUp(self):
        self.superuser = User.objects.create_superuser('admin_mob', 'admin_mob@example.com', 'pass1234')
        self.user = User.objects.create_user('emp_mob', 'emp_mob@example.com', 'pass1234')
        self.employee = Employee.objects.create(
            user=self.user, name='Empleado Móvil', email='emp_mob@example.com', hire_date=date(2026, 1, 1)
        )
        Salary.objects.create(
            employee=self.employee, base_amount=Decimal('600.00'), effective_date=date(2026, 1, 1)
        )

    def test_navbar_has_mobile_toggler_and_responsive_tables(self):
        self.client.login(username='emp_mob', password='pass1234')
        resp = self.client.get(reverse('mi_panel'))
        self.assertEqual(resp.status_code, 200)
        html = resp.content.decode('utf-8')
        self.assertIn('class="navbar-toggler"', html)
        self.assertIn('data-target="#navbarNav"', html)
        self.assertIn('table-responsive', html)

    def test_admin_base_site_has_mobile_viewport_and_responsive_styles(self):
        self.client.login(username='admin_mob', password='pass1234')
        resp = self.client.get(reverse('admin:index'))
        self.assertEqual(resp.status_code, 200)
        html = resp.content.decode('utf-8')
        self.assertIn('name="viewport"', html)
        self.assertIn('@media (max-width: 767px)', html)
        self.assertIn('Dashboard', html)
